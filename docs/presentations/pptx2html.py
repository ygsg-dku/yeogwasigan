"""생성된 .pptx를 읽어 HTML로 미러링한다. 시각 판정은 소스가 아니라 산출물을 본다."""
import sys, html
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

PX = 96.0  # 1in = 96px


def rgb(c):
    try:
        return "#%02X%02X%02X" % (c[0], c[1], c[2])
    except Exception:
        return None


def solid(fill):
    try:
        if fill.type is not None and int(fill.type) == 1:
            return rgb(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def emit(path, out):
    prs = Presentation(path)
    W = Emu(prs.slide_width).inches * PX
    H = Emu(prs.slide_height).inches * PX
    parts = ["<style>@page{size:%dpx %dpx;margin:0}body{margin:0}"
             ".s{position:relative;overflow:hidden;page-break-after:always}"
             ".sh{position:absolute;box-sizing:border-box}"
             ".t{position:absolute;box-sizing:border-box;white-space:pre-wrap;"
             "line-height:1.25;font-family:'Arial Unicode MS',sans-serif}</style>" % (W, H)]
    for s in prs.slides:
        bg = solid(s.background.fill) or "#FFFFFF"
        parts.append(f'<div class="s" style="width:{W}px;height:{H}px;background:{bg}">')
        for sh in s.shapes:
            x = Emu(sh.left).inches * PX
            y = Emu(sh.top).inches * PX
            w = Emu(sh.width).inches * PX
            h = Emu(sh.height).inches * PX
            if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                f = solid(sh.fill) or "transparent"
                try:
                    if int(sh.fill.type) == 2:  # patterned
                        fg = rgb(sh.fill.fore_color.rgb) or "#000"
                        bg = rgb(sh.fill.back_color.rgb) or "#FFF"
                        f = (f"repeating-linear-gradient(45deg,{fg} 0 2px,"
                             f"{bg} 2px 4px)")
                except Exception:
                    pass
                try:
                    ln = rgb(sh.line.color.rgb)
                    lw = Emu(sh.line.width).inches * PX if sh.line.width else 1
                except Exception:
                    ln, lw = None, 0
                brd = f"border:{lw:.1f}px solid {ln};" if ln else ""
                nm = str(sh.auto_shape_type)
                rad = "border-radius:50%;" if "OVAL" in nm else (
                    "border-radius:7px;" if "ROUNDED" in nm else "")
                parts.append(f'<div class="sh" style="left:{x:.1f}px;top:{y:.1f}px;'
                             f'width:{w:.1f}px;height:{h:.1f}px;background:{f};{brd}{rad}"></div>')
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            va = {2: "flex-end", 3: "center"}.get(
                int(tf.vertical_anchor) if tf.vertical_anchor is not None else 1, "flex-start")
            inner = []
            for p in tf.paragraphs:
                if not "".join(r.text for r in p.runs).strip():
                    continue
                al = {2: "center", 3: "right"}.get(
                    int(p.alignment) if p.alignment is not None else 1, "left")
                spans = []
                for r in p.runs:
                    sz = (r.font.size.pt if r.font.size else 14) * PX / 72.0
                    try:
                        col = rgb(r.font.color.rgb) or "#000"
                    except Exception:
                        col = "#000"
                    bd = "font-weight:700;" if r.font.bold else ""
                    spans.append(f'<span style="font-size:{sz:.1f}px;color:{col};{bd}">'
                                 f'{html.escape(r.text)}</span>')
                inner.append(f'<div style="text-align:{al}">' + "".join(spans) + "</div>")
            parts.append(f'<div class="t" style="left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;'
                         f'height:{h:.1f}px;display:flex;flex-direction:column;'
                         f'justify-content:{va}">' + "".join(inner) + "</div>")
        parts.append("</div>")
    open(out, "w", encoding="utf-8").write("".join(parts))
    print("wrote", out)


if __name__ == "__main__":
    emit(sys.argv[1], sys.argv[2])
